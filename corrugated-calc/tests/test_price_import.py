"""Импорт на ценоразписи: всички формати, преглед, редакция, валидиране, одобрение, активиране."""

from decimal import Decimal as D

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, text

import fixtures_gen as g

BOX = {"length_mm": 1192, "width_mm": 792, "height_mm": 779}
MGR = {"X-User": "manager"}
ADMIN = {"X-User": "admin"}


@pytest.fixture()
def client(db):
    from calc.api.app import app
    return TestClient(app)


def upload(client, files, code="NEW-SUPPLIER-2026", supplier="Нов доставчик", user=MGR):
    r = client.post("/api/v1/imports", headers=user, data={"price_list_code": code, "supplier": supplier},
                    files=[("files", (name, data)) for name, data in files])
    assert r.status_code == 200, r.text
    return r.json()


def edits(client, job_id, ops, user=MGR):
    r = client.post(f"/api/v1/imports/{job_id}/edits", headers=user, json={"ops": ops})
    assert r.status_code == 200, r.text
    return r.json()


def validate(client, job_id):
    r = client.post(f"/api/v1/imports/{job_id}/validate", headers=MGR)
    assert r.status_code == 200, r.text
    return r.json()


def codes(issues):
    return {i["code"] for i in issues}


def calc_total(client, qty=160, code="EC3008"):
    r = client.post("/api/v1/calculate", json={"box": BOX, "material_code": code, "quantity": qty})
    return r.json()["material"]


FORMATS = [
    ("cennik.xlsx", g.xlsx, "sheet", "excel"),
    ("cennik.xls", g.xls, "sheet", "excel"),
    ("cennik.csv", g.csv_bytes, "sheet", "csv"),
    ("cennik.pdf", g.text_pdf, "page", "pdf_text"),
    ("scan.pdf", g.scanned_pdf, "page", "ocr"),
    ("snimka.png", g.image, "image", "ocr"),
    ("snimka.jpg", lambda: g.image(fmt="JPEG"), "image", "ocr"),
    ("snimka.heic", lambda: g.image(fmt="HEIC"), "image", "ocr"),
    ("zavartyana.jpg", lambda: g.image(fmt="JPEG", rotate_exif=True), "image", "ocr"),
]


@pytest.mark.parametrize("name,make,kind,method", FORMATS)
def test_full_flow_each_format(client, name, make, kind, method):
    job = upload(client, [(name, make())])
    p = job["proposal"]
    assert job["status"] == "extracted"
    assert {(t["lower"]["value"], t["upper"]["value"]) for t in p["tiers"]} == {("1000", "1500"), ("1500", "2000")}
    rows = {r["material"]["value"]: r for r in p["rows"]}
    assert set(rows) == {"EC3008", "EC4027"}
    assert D(rows["EC3008"]["prices"]["T1"]["value"]) == D("1060")
    assert D(rows["EC3008"]["prices"]["T2"]["value"]) == D("963")
    assert D(rows["EC4027"]["prices"]["T1"]["value"]) == D("1003")
    assert "T2" not in rows["EC4027"]["prices"] or rows["EC4027"]["prices"]["T2"]["value"] is None
    assert (p["currency"]["value"], p["price_unit"]["value"], p["tax_basis"]["value"]) == \
        ("BGN", "per_1000m2", "excl_vat")
    assert p["valid_from"]["value"] == "2026-03-01" and p["min_order_m2"]["value"] == "1000"

    # връзка към източника за всяка клетка
    src = rows["EC3008"]["prices"]["T1"]["source"]
    assert src["file_id"] == job["files"][0]["id"] and src["kind"] == kind
    if kind == "sheet":
        assert src["cell"] and src["sheet"]
    else:
        assert len(src["bbox"]) == 4
        if kind == "page":
            assert src["page"] == 1
    assert any(t["method"] == method for t in job["extraction"]["tables"])

    # не се активира автоматично: калкулацията ползва стария (reference) ценоразпис
    assert calc_total(client)["source"]["price_list_id"] == "TEST-SUPPLIER-2026"

    v = validate(client, job["id"])
    assert not v["ok"] and "TIER_BOUNDS_UNCONFIRMED" in codes(v["errors"])
    edits(client, job["id"], [{"op": "confirm_tiers"}])
    v = validate(client, job["id"])
    assert v["ok"], v["errors"]
    r = client.post(f"/api/v1/imports/{job['id']}/approve", headers=MGR, json={"acknowledge_warnings": True})
    assert r.status_code == 200, r.text
    assert r.json() == {"price_list_code": "NEW-SUPPLIER-2026", "version": 1, "status": "approved"}
    assert calc_total(client)["source"]["price_list_id"] == "TEST-SUPPLIER-2026"  # approved ≠ active

    r = client.post("/api/v1/price-lists/NEW-SUPPLIER-2026/versions/1/activate", headers=ADMIN)
    assert r.status_code == 200, r.text
    m = calc_total(client)
    assert m["source"]["price_list_id"] == "NEW-SUPPLIER-2026" and m["source"]["price_list_status"] == "active"
    assert m["total"]["amount"] == "1067.8979328"

    ver = client.get("/api/v1/price-lists/NEW-SUPPLIER-2026/versions/1", headers=MGR).json()
    assert ver["prices"][0]["source_ref"]["import_id"] == job["id"]


def test_previews_and_original_kept(client, db):
    heic = g.image(fmt="HEIC")
    job = upload(client, [("a.heic", heic), ("b.pdf", g.text_pdf())])
    f_img, f_pdf = job["files"]
    r = client.get(f"/api/v1/imports/{job['id']}/files/{f_img['id']}/preview", headers=MGR)
    assert r.status_code == 200 and r.content[:4] == b"\x89PNG"
    r = client.get(f"/api/v1/imports/{job['id']}/files/{f_pdf['id']}/preview?page=1", headers=MGR)
    assert r.status_code == 200 and r.content[:4] == b"\x89PNG"
    r = client.get(f"/api/v1/imports/{job['id']}/files/{f_img['id']}", headers=MGR)
    assert r.content == heic and r.headers["content-security-policy"] == "sandbox"
    with create_engine(db).begin() as c, pytest.raises(Exception, match="Оригиналният файл не се променя"):
        c.execute(text("update import_files set content = 'x' where id = :i"), {"i": f_img["id"]})


def test_multiple_images_one_price_list(client):
    page1 = g.image(table=[g.TABLE[0], g.TABLE[1]])
    page2 = g.image(table=[g.TABLE[2]], header_lines=[])
    job = upload(client, [("p1.jpg", page1), ("p2.jpg", page2)])
    rows = {r["material"]["value"]: r for r in job["proposal"]["rows"]}
    assert set(rows) == {"EC3008", "EC4027"}
    assert rows["EC4027"]["prices"]["T1"]["source"]["file_id"] == job["files"][1]["id"]
    assert any("предходната таблица" in n for n in job["proposal"]["notes"])


def test_ambiguous_and_unreadable_not_guessed(client):
    table = [g.TABLE[0], ["EC3008", "1,060", "963"], ["EC4027", "1O03", ""]]
    job = upload(client, [("x.csv", g.csv_bytes(table))])
    rows = {r["material"]["value"]: r for r in job["proposal"]["rows"]}
    amb = rows["EC3008"]["prices"]["T1"]
    assert amb["status"] == "ambiguous" and amb["value"] is None and set(amb["candidates"]) == {"1060", "1.060"}
    bad = rows["EC4027"]["prices"]["T1"]
    assert bad["status"] == "unreadable" and bad["value"] is None
    edits(client, job["id"], [{"op": "confirm_tiers"}])
    v = validate(client, job["id"])
    assert "PRICE_REVIEW" in codes(v["errors"])
    r1, r2 = rows["EC3008"]["id"], rows["EC4027"]["id"]
    # двусмислено въведено от потребителя също се отказва
    r = client.post(f"/api/v1/imports/{job['id']}/edits", headers=MGR,
                    json={"ops": [{"op": "set_price", "row_id": r1, "tier": "T1", "value": "1,060"}]})
    assert r.status_code == 422
    edits(client, job["id"], [{"op": "set_price", "row_id": r1, "tier": "T1", "value": "1060"},
                              {"op": "set_price", "row_id": r2, "tier": "T1", "value": "1003"}])
    assert validate(client, job["id"])["ok"]


def test_decimal_separator_override(client):
    table = [g.TABLE[0], ["EC3008", "1.060", "963"], ["EC4027", "1.003", ""]]
    job = upload(client, [("x.csv", g.csv_bytes(table))])
    assert job["proposal"]["rows"][0]["prices"]["T1"]["status"] == "ambiguous"
    r = client.post(f"/api/v1/imports/{job['id']}/remap", headers=MGR, json={"decimal_separator": "comma"})
    p = r.json()["proposal"]
    assert p["rows"][0]["prices"]["T1"]["value"] == "1060" and p["decimal_separator"] == "comma"


def test_excel_sheet_selection_and_manual_mapping(client):
    import io
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(g.xlsx()))
    other = wb.create_sheet("Бележки", 0)
    other.append(["Вътрешни бележки"])
    raw = wb.create_sheet("Без заглавие")
    raw.append(["EC3008", None, 1100, 990])
    buf = io.BytesIO()
    wb.save(buf)
    job = upload(client, [("multi.xlsx", buf.getvalue())])
    labels = {t["label"] for t in job["extraction"]["tables"]}
    assert labels == {"Бележки", "Цени", "Без заглавие"}
    maps = {m["table_id"].split(":", 1)[1]: m for m in job["proposal"]["mappings"]}
    assert maps["Бележки"]["skip"]
    # листът без заглавие се третира като продължение → дублиран EC3008
    v = validate(client, edits(client, job["id"], [{"op": "confirm_tiers"}])["id"])
    assert "DUPLICATE_MATERIAL" in codes(v["errors"])
    # ръчно: изключваме листа
    tid = next(t["id"] for t in job["extraction"]["tables"] if t["label"] == "Без заглавие")
    cid = next(t["id"] for t in job["extraction"]["tables"] if t["label"] == "Цени")
    r = client.post(f"/api/v1/imports/{job['id']}/remap", headers=MGR, json={
        "mappings": [{"table_id": tid, "skip": True}, maps["Цени"] | {"table_id": cid}], "keep_tiers": True})
    assert len(r.json()["proposal"]["rows"]) == 2


def test_validation_errors(client):
    job = upload(client, [("x.csv", g.csv_bytes([g.TABLE[0], ["EC9999", "1000", "900"]]))], code=None,
                 supplier=None)
    edits(client, job["id"], [{"op": "set_field", "field": "currency", "value": None},
                              {"op": "set_tier", "code": "T2", "lower": "1400", "upper": "2000", "confirm": True},
                              {"op": "confirm_tiers"}])
    v = validate(client, job["id"])
    c = codes(v["errors"])
    assert {"UNKNOWN_MATERIAL", "TIER_OVERLAP", "MISSING"} <= c
    assert not client.post(f"/api/v1/imports/{job['id']}/approve", headers=MGR, json={}).status_code == 200


def _activate_v1(client):
    job = upload(client, [("v1.xlsx", g.xlsx())])
    edits(client, job["id"], [{"op": "confirm_tiers"}])
    assert validate(client, job["id"])["ok"]
    client.post(f"/api/v1/imports/{job['id']}/approve", headers=MGR, json={"acknowledge_warnings": True})
    assert client.post("/api/v1/price-lists/NEW-SUPPLIER-2026/versions/1/activate", headers=ADMIN).status_code == 200


def test_unusual_change_versioning_and_immutability(client, db):
    _activate_v1(client)
    new_table = g.table_with({"EC3008": ("1272,00", "970,00"), "EC4027": ("1010,00", "")})  # +20% за T1
    job = upload(client, [("v2.xlsx", g.xlsx(new_table))])
    edits(client, job["id"], [{"op": "confirm_tiers"}])
    v = validate(client, job["id"])
    assert v["ok"]
    unusual = [w for w in v["warnings"] if w["code"] == "UNUSUAL_CHANGE"]
    assert len(unusual) == 1 and unusual[0]["ref"]["old"] == "1060" and unusual[0]["ref"]["pct"] == "20.00"
    r = client.post(f"/api/v1/imports/{job['id']}/approve", headers=MGR, json={})
    assert r.status_code == 422 and "потвърдете" in r.json()["detail"]
    r = client.post(f"/api/v1/imports/{job['id']}/approve", headers=MGR, json={"acknowledge_warnings": True})
    assert r.json()["version"] == 2

    before = calc_total(client)
    assert before["source"]["price_list_version"] == 1
    client.post("/api/v1/price-lists/NEW-SUPPLIER-2026/versions/2/activate", headers=ADMIN)
    after = calc_total(client)
    assert after["source"]["price_list_version"] == 2 and after["tier"]["price"] == "1272"

    lists = client.get("/api/v1/price-lists", headers=MGR).json()
    vers = {v["version"]: v["status"] for v in next(p for p in lists if p["code"] == "NEW-SUPPLIER-2026")["versions"]}
    assert vers == {1: "archived", 2: "active"}
    v1 = client.get("/api/v1/price-lists/NEW-SUPPLIER-2026/versions/1", headers=MGR).json()
    assert {p["price"] for p in v1["prices"] if p["material_code"] == "EC3008"} == {"1060", "963"}

    with create_engine(db).begin() as c, pytest.raises(Exception, match="не се променят"):
        c.execute(text("update material_price_tiers set price = 1 where version_id = "
                       "(select id from price_list_versions where version = 1 and price_list_id = "
                       "(select id from price_lists where code = 'NEW-SUPPLIER-2026'))"))
    with create_engine(db).begin() as c, pytest.raises(Exception, match="Недопустим преход"):
        c.execute(text("update price_list_versions set status = 'active' where version = 1 and price_list_id = "
                       "(select id from price_lists where code = 'NEW-SUPPLIER-2026')"))
    with create_engine(db).begin() as c, pytest.raises(Exception, match="само за добавяне"):
        c.execute(text("delete from audit_log"))


def test_price_edit_is_audited(client):
    job = upload(client, [("x.xlsx", g.xlsx())])
    row = job["proposal"]["rows"][0]
    edits(client, job["id"], [{"op": "set_price", "row_id": row["id"], "tier": "T1", "value": "1075.50"}])
    log = client.get(f"/api/v1/audit?entity=import_job&entity_id={job['id']}", headers=ADMIN).json()
    e = next(a for a in log if a["action"] == "import.edit.set_price")
    assert e["actor_id"] == client.get("/api/v1/me", headers=MGR).json()["id"]
    assert e["at"] and e["old"]["value"] == "1060" and e["new"]["value"] == "1075.50"
    assert e["new"]["source"]["cell"]  # редакцията пази връзката към оригинала


def test_edit_after_validation_requires_revalidation(client):
    job = upload(client, [("x.xlsx", g.xlsx())])
    edits(client, job["id"], [{"op": "confirm_tiers"}])
    assert validate(client, job["id"])["ok"]
    edits(client, job["id"], [{"op": "set_field", "field": "valid_to", "value": "2026-11-30"}])
    r = client.post(f"/api/v1/imports/{job['id']}/approve", headers=MGR, json={"acknowledge_warnings": True})
    assert r.status_code == 422


def test_roles(client):
    r = client.post("/api/v1/imports", headers={"X-User": "sales"}, files=[("files", ("x.csv", g.csv_bytes()))])
    assert r.status_code == 403
    assert client.post("/api/v1/imports", files=[("files", ("x.csv", g.csv_bytes()))]).status_code == 401
    assert client.get("/api/v1/imports", headers={"X-User": "nobody"}).status_code == 401
    r = client.post("/api/v1/price-lists/TEST-SUPPLIER-2026/versions/1/activate", headers=MGR)
    assert r.status_code == 403


def test_manual_entry_and_draft_from_version(client):
    r = client.post("/api/v1/imports/manual", headers=MGR, json={"price_list_code": "MANUAL-1", "supplier": "Ръчен"})
    job = r.json()
    ops = [{"op": "add_tier", "lower": "1000", "upper": "1500", "confirm": True},
           {"op": "add_tier", "lower": "1500", "confirm": True},
           {"op": "add_row", "material": "EC3008", "prices": {"T1": "1050", "T2": "950"}},
           {"op": "set_field", "field": "currency", "value": "BGN"},
           {"op": "set_field", "field": "price_unit", "value": "per_1000m2"},
           {"op": "set_field", "field": "tax_basis", "value": "excl_vat"},
           {"op": "set_field", "field": "valid_from", "value": "2026-01-01"}]
    edits(client, job["id"], ops)
    assert validate(client, job["id"])["ok"]
    assert client.post(f"/api/v1/imports/{job['id']}/approve", headers=MGR, json={}).json()["version"] == 1

    d = client.post("/api/v1/price-lists/MANUAL-1/versions/1/draft", headers=MGR).json()
    row = d["proposal"]["rows"][0]
    edits(client, d["id"], [{"op": "set_price", "row_id": row["id"], "tier": "T1", "value": "1055"},
                            {"op": "set_field", "field": "valid_from", "value": "2026-06-01"}])
    assert validate(client, d["id"])["ok"]
    assert client.post(f"/api/v1/imports/{d['id']}/approve", headers=MGR, json={}).json()["version"] == 2
    v1 = client.get("/api/v1/price-lists/MANUAL-1/versions/1", headers=MGR).json()
    assert {p["price"] for p in v1["prices"]} == {"1050", "950"}


def test_corrupt_file_allows_manual_entry(client):
    job = upload(client, [("broken.xlsx", b"PK\x03\x04garbage")])
    assert job["extraction"]["errors"] and "ръчно" in job["extraction"]["errors"][0]["message"]
    assert job["proposal"]["rows"] == []


def test_seeded_db_calculation_matches_acceptance(client):
    m = calc_total(client, 238)
    assert m["total"]["amount"] == "1588.49817504"
    o = client.post("/api/v1/calculate", json={"box": BOX, "material_code": "EC3008", "quantity": 238}).json()
    assert o["optimization"]["proposed_total"]["amount"] == "1449.199201176"


def test_implausible_value_warns(client):
    job = upload(client, [("x.csv", g.csv_bytes([g.TABLE[0], ["EC3008", "1,06", "963"], ["EC4027", "1003", ""]]))])
    edits(client, job["id"], [{"op": "confirm_tiers"}])
    v = validate(client, job["id"])
    assert v["ok"]
    c = codes(v["warnings"])
    assert {"PRICE_OUTLIER", "TIER_PRICE_INCREASES"} <= c
    assert client.post(f"/api/v1/imports/{job['id']}/approve", headers=MGR, json={}).status_code == 422


def test_malformed_ops_rejected(client):
    job = upload(client, [("x.csv", g.csv_bytes())])
    for op in [{"op": "set_price"}, {"op": "confirm", "field": "__class__"}, {"op": "nope"},
               {"op": "set_price", "row_id": "r999", "tier": "T1", "value": "1"}]:
        r = client.post(f"/api/v1/imports/{job['id']}/edits", headers=MGR, json={"ops": [op]})
        assert r.status_code == 422, op
