"""Извличане на таблици от Excel (.xlsx/.xls), CSV, текстов PDF, сканиран PDF и снимки.

Всяка клетка пази връзка към източника: файл + sheet/ред/колона или страница/снимка + bbox.
bbox е в пиксели на preview изображението (PDF страници се рендират с PREVIEW_SCALE).
Резултатът е предложение за преглед; нищо не се активира автоматично.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from typing import Any

PREVIEW_SCALE = 2.0  # 144 dpi за PDF preview
OCR_SCALE = 300 / 72  # 300 dpi за OCR на сканиран PDF
OCR_LANG = "bul+eng"
LOW_CONFIDENCE = 70

IMAGE_EXT = {".jpg", ".jpeg", ".png", ".heic", ".heif", ".webp", ".tif", ".tiff"}


class ExtractionError(ValueError):
    pass


@dataclass
class FileInfo:
    id: int
    filename: str
    content: bytes


def detect_kind(filename: str, content: bytes) -> str:
    name = filename.lower()
    head = content[:8]
    if head.startswith(b"%PDF"):
        return "pdf"
    if head.startswith(b"PK") and name.endswith((".xlsx", ".xlsm")):
        return "xlsx"
    if head.startswith(b"\xd0\xcf\x11\xe0"):
        return "xls"
    if any(name.endswith(e) for e in IMAGE_EXT) or head.startswith((b"\x89PNG", b"\xff\xd8\xff")) \
            or content[4:12] in (b"ftypheic", b"ftypheix", b"ftypmif1", b"ftyphevc"):
        return "image"
    if name.endswith((".csv", ".txt", ".tsv")):
        return "csv"
    raise ExtractionError(f"Неподдържан формат: {filename}")


def col_letter(c: int) -> str:
    s = ""
    c += 1
    while c:
        c, r = divmod(c - 1, 26)
        s = chr(65 + r) + s
    return s


def _cell(text: str, r: int, c: int, **kw) -> dict[str, Any]:
    d = {"t": text, "r": r, "c": c}
    d.update({k: v for k, v in kw.items() if v is not None})
    return d


def _table(file: FileInfo, kind: str, label: str, method: str, rows: list[list[dict]], **kw) -> dict[str, Any]:
    t = {"id": f"f{file.id}:{label}", "file_id": file.id, "filename": file.filename, "kind": kind, "label": label,
         "method": method, "rows": rows}
    t.update({k: v for k, v in kw.items() if v is not None})
    return t


# ---------------------------------------------------------------- Excel / CSV

def _num_text(v) -> tuple[str, bool]:
    """Числова клетка от Excel → текст без загуба (repr на double е най-краткото точно представяне)."""
    if isinstance(v, bool):
        return str(v), False
    if isinstance(v, int):
        return str(v), True
    if isinstance(v, float):
        return (str(int(v)) if v.is_integer() else repr(v)), True
    if hasattr(v, "isoformat"):
        return v.isoformat()[:10], False
    return str(v), False


def extract_xlsx(file: FileInfo) -> list[dict]:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(file.content), data_only=True, read_only=True)
    tables = []
    for ws in wb.worksheets:
        rows = []
        for r, row in enumerate(ws.iter_rows(values_only=True)):
            cells = []
            for c, v in enumerate(row):
                if v is None or (isinstance(v, str) and not v.strip()):
                    continue
                text, is_num = _num_text(v)
                cells.append(_cell(text.strip(), r, c, cell=f"{col_letter(c)}{r + 1}", num=is_num or None))
            rows.append(cells)
        tables.append(_table(file, "sheet", ws.title, "excel", _trim(rows), sheet=ws.title))
    return tables


def extract_xls(file: FileInfo) -> list[dict]:
    import xlrd
    book = xlrd.open_workbook(file_contents=file.content)
    tables = []
    for sh in book.sheets():
        rows = []
        for r in range(sh.nrows):
            cells = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    continue
                if cell.ctype == xlrd.XL_CELL_NUMBER:
                    text, is_num = _num_text(cell.value)
                elif cell.ctype == xlrd.XL_CELL_DATE:
                    text, is_num = xlrd.xldate.xldate_as_datetime(cell.value, book.datemode).date().isoformat(), False
                else:
                    text, is_num = str(cell.value).strip(), False
                if text:
                    cells.append(_cell(text, r, c, cell=f"{col_letter(c)}{r + 1}", num=is_num or None))
            rows.append(cells)
        tables.append(_table(file, "sheet", sh.name, "excel", _trim(rows), sheet=sh.name))
    return tables


def extract_csv(file: FileInfo) -> list[dict]:
    for enc in ("utf-8-sig", "cp1251"):
        try:
            text = file.content.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ExtractionError("Неразпознато кодиране на CSV")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t|")
    except csv.Error:
        dialect = csv.excel
        dialect.delimiter = ";" if text.count(";") > text.count(",") else ","
    rows = []
    for r, row in enumerate(csv.reader(io.StringIO(text), dialect)):
        rows.append([_cell(v.strip(), r, c, cell=f"{col_letter(c)}{r + 1}") for c, v in enumerate(row) if v.strip()])
    return [_table(file, "sheet", "csv", "csv", _trim(rows), sheet="csv")]


def _trim(rows: list[list[dict]]) -> list[list[dict]]:
    while rows and not rows[-1]:
        rows.pop()
    return rows


# ---------------------------------------------------------------- думи → редове/колони

def words_to_grid(words: list[dict]) -> list[list[dict]]:
    """words: {text, x0, x1, top, bottom, conf?}. Групира в редове и колони по проекция на интервалите."""
    words = [w for w in words if w["text"].strip()]
    if not words:
        return []
    words.sort(key=lambda w: (w["top"] + w["bottom"]) / 2)
    lines: list[list[dict]] = []
    for w in words:
        mid = (w["top"] + w["bottom"]) / 2
        h = w["bottom"] - w["top"]
        if lines:
            last = lines[-1]
            lmid = sum((x["top"] + x["bottom"]) / 2 for x in last) / len(last)
            lh = max(x["bottom"] - x["top"] for x in last)
            if abs(mid - lmid) < 0.55 * max(h, lh):
                last.append(w)
                continue
        lines.append([w])

    phrases_per_line = []
    for line in lines:
        line.sort(key=lambda w: w["x0"])
        h = max(w["bottom"] - w["top"] for w in line)
        phrases, cur = [], [line[0]]
        for w in line[1:]:
            if w["x0"] - cur[-1]["x1"] < 0.9 * h:
                cur.append(w)
            else:
                phrases.append(cur)
                cur = [w]
        phrases.append(cur)
        phrases_per_line.append([{
            "text": " ".join(w["text"] for w in p), "x0": min(w["x0"] for w in p), "x1": max(w["x1"] for w in p),
            "top": min(w["top"] for w in p), "bottom": max(w["bottom"] for w in p),
            "conf": min((w.get("conf", 100) for w in p), default=100)} for p in phrases])

    bands: list[list[float]] = []
    for ph in phrases_per_line:
        if len(ph) < 2:
            continue
        for p in ph:
            bands.append([p["x0"], p["x1"]])
    bands.sort()
    merged: list[list[float]] = []
    for b in bands:
        if merged and b[0] <= merged[-1][1]:
            merged[-1][1] = max(merged[-1][1], b[1])
        else:
            merged.append(list(b))

    def col_of(p):
        if not merged:
            return 0
        best, best_ov = 0, -1e9
        for i, (a, b) in enumerate(merged):
            ov = min(b, p["x1"]) - max(a, p["x0"])
            if ov > best_ov:
                best, best_ov = i, ov
        return best

    grid = []
    for r, ph in enumerate(phrases_per_line):
        cells: dict[int, dict] = {}
        for p in ph:
            c = col_of(p)
            bbox = [round(p["x0"], 1), round(p["top"], 1), round(p["x1"], 1), round(p["bottom"], 1)]
            if c in cells:  # две фрази в една колона → слепване
                prev = cells[c]
                prev["t"] += " " + p["text"]
                prev["bbox"] = [min(prev["bbox"][0], bbox[0]), min(prev["bbox"][1], bbox[1]),
                                max(prev["bbox"][2], bbox[2]), max(prev["bbox"][3], bbox[3])]
                prev["conf"] = min(prev.get("conf", 100), p["conf"])
            else:
                cells[c] = _cell(p["text"], r, c, bbox=bbox, conf=round(p["conf"]) if p["conf"] < 100 else None)
        grid.append([cells[c] for c in sorted(cells)])
    return grid


# ---------------------------------------------------------------- PDF

def extract_pdf(file: FileInfo) -> list[dict]:
    import pdfplumber
    tables = []
    with pdfplumber.open(io.BytesIO(file.content)) as pdf:
        for pno, page in enumerate(pdf.pages, start=1):
            size = [round(float(page.width) * PREVIEW_SCALE), round(float(page.height) * PREVIEW_SCALE)]
            if len(page.chars) >= 10:
                tables.extend(_pdf_text_page(file, page, pno, size))
            else:
                tables.append(_ocr_pdf_page(file, pno, size))
    return tables


def _scale(b, k=PREVIEW_SCALE):
    return [round(float(x) * k, 1) for x in b]


def _pdf_text_page(file: FileInfo, page, pno: int, size) -> list[dict]:
    out = []
    found = page.find_tables()
    for ti, tb in enumerate(found):
        texts = tb.extract()
        rows = []
        for r, row in enumerate(tb.rows):
            cells = []
            for c, bbox in enumerate(row.cells):
                txt = texts[r][c] if r < len(texts) and c < len(texts[r]) else None
                if bbox is None or not txt or not str(txt).strip():
                    continue
                cells.append(_cell(" ".join(str(txt).split()), r, c, bbox=_scale(bbox)))
            rows.append(cells)
        label = f"стр. {pno}" + (f" т.{ti + 1}" if len(found) > 1 else "")
        out.append(_table(file, "page", label, "pdf_text", _trim(rows), page=pno, image_size=size))
    def inside(w, b):
        return b[0] - 1 <= float(w["x0"]) and float(w["x1"]) <= b[2] + 1 and b[1] - 1 <= float(w["top"]) \
            and float(w["bottom"]) <= b[3] + 1

    all_words = page.extract_words(keep_blank_chars=False, use_text_flow=False)
    free = [w for w in all_words if not any(inside(w, tb.bbox) for tb in found)]
    words = [{"text": w["text"], "x0": float(w["x0"]) * PREVIEW_SCALE, "x1": float(w["x1"]) * PREVIEW_SCALE,
              "top": float(w["top"]) * PREVIEW_SCALE, "bottom": float(w["bottom"]) * PREVIEW_SCALE} for w in free]
    if not out:
        out.append(_table(file, "page", f"стр. {pno}", "pdf_text", words_to_grid(words), page=pno, image_size=size))
    elif words:
        # текст извън таблиците (валута, ДДС, валидност) — само за документните полета, не за редове
        out.insert(0, _table(file, "text", f"стр. {pno} текст", "pdf_text", words_to_grid(words), page=pno,
                             image_size=size))
    return out


def render_pdf_page(content: bytes, pno: int, scale: float = PREVIEW_SCALE):
    import pypdfium2 as pdfium
    pdf = pdfium.PdfDocument(content)
    try:
        return pdf[pno - 1].render(scale=scale).to_pil()
    finally:
        pdf.close()


def pdf_page_count(content: bytes) -> int:
    import pypdfium2 as pdfium
    pdf = pdfium.PdfDocument(content)
    try:
        return len(pdf)
    finally:
        pdf.close()


def _ocr_pdf_page(file: FileInfo, pno: int, size) -> dict:
    img = render_pdf_page(file.content, pno, OCR_SCALE)
    grid = ocr_grid(img, back_scale=PREVIEW_SCALE / OCR_SCALE)
    return _table(file, "page", f"стр. {pno} (OCR)", "ocr", grid, page=pno, image_size=size)


# ---------------------------------------------------------------- снимки и OCR

def load_image(content: bytes):
    from PIL import Image, ImageOps
    try:
        import pillow_heif
        pillow_heif.register_heif_opener()
    except ImportError:  # pragma: no cover
        pass
    img = Image.open(io.BytesIO(content))
    img = ImageOps.exif_transpose(img)
    return img.convert("RGB")


def image_preview_png(content: bytes) -> bytes:
    """Безопасно преобразуване (вкл. HEIC и EXIF ориентация) в PNG за преглед."""
    buf = io.BytesIO()
    load_image(content).save(buf, format="PNG")
    return buf.getvalue()


def ocr_grid(img, back_scale: float = 1.0) -> list[list[dict]]:
    import pytesseract
    from PIL import ImageOps
    gray = ImageOps.grayscale(img)
    up = 1.0
    if gray.width < 1600:
        up = 2.0
        gray = gray.resize((gray.width * 2, gray.height * 2))
    data = pytesseract.image_to_data(gray, lang=OCR_LANG, config="--psm 6 -c preserve_interword_spaces=1",
                                     output_type=pytesseract.Output.DICT)
    k = back_scale / up
    words = []
    for i, txt in enumerate(data["text"]):
        conf = float(data["conf"][i])
        if conf < 0 or not txt.strip():
            continue
        x, y, w, h = data["left"][i], data["top"][i], data["width"][i], data["height"][i]
        words.append({"text": txt, "x0": x * k, "x1": (x + w) * k, "top": y * k, "bottom": (y + h) * k,
                      "conf": conf})
    return words_to_grid(words)


def extract_image(file: FileInfo, index: int) -> list[dict]:
    img = load_image(file.content)
    return [_table(file, "image", f"снимка {index}", "ocr", ocr_grid(img), image_size=[img.width, img.height])]


# ---------------------------------------------------------------- вход

def extract_files(files: list[FileInfo]) -> dict[str, Any]:
    tables, errors = [], []
    img_no = 0
    for f in files:
        try:
            kind = detect_kind(f.filename, f.content)
            if kind == "xlsx":
                tables += extract_xlsx(f)
            elif kind == "xls":
                tables += extract_xls(f)
            elif kind == "csv":
                tables += extract_csv(f)
            elif kind == "pdf":
                tables += extract_pdf(f)
            else:
                img_no += 1
                tables += extract_image(f, img_no)
        except ExtractionError as e:
            errors.append({"file_id": f.id, "filename": f.filename, "message": str(e)})
        except Exception as e:  # повреден файл и др.: не спира останалите
            errors.append({"file_id": f.id, "filename": f.filename,
                           "message": f"Файлът не може да бъде разчетен автоматично ({type(e).__name__}: {e}). "
                                      f"Въведете данните ръчно по оригинала."})
    return {"tables": tables, "errors": errors}


def text_lines(extraction: dict[str, Any]) -> list[tuple[str, dict]]:
    """Всички редове като текст с източник — за валута, ДДС, валидност, минимална заявка."""
    out = []
    for t in extraction.get("tables", []):
        for row in t["rows"]:
            if row:
                out.append((" ".join(c["t"] for c in row), {"table_id": t["id"], "row": row[0]["r"],
                                                            "cell": row[0]}))
    return out


WS = re.compile(r"\s+")
